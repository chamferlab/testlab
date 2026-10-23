"""PowerPoint slides of the Final Report's runs: an editable copy from the frozen records."""

import io

import pytest
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from PIL import Image
from pptx import Presentation

from testlab.documents import chromium
from testlab.projects import drawings
from testlab.reports import slides
from testlab.results import finalization, images
from testlab.results.models import ResultImage, TestResult
from testlab.results.tests.test_finalization import complete_result

from .test_final_report import include, second_run

pytestmark = pytest.mark.django_db


def png(size=(80, 60), color="white"):
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, "PNG")
    return out.getvalue()


def add_picture(result, user, category, label=None, caption=""):
    component = next((c for c in result.components.all() if c.label == label), None) if label else None
    images.add_image(result, SimpleUploadedFile("photo.png", png()), user, category, component, caption)


@pytest.fixture(autouse=True)
def no_local_template(settings, tmp_path):
    """A template in this machine's data/ folder never changes the tests."""
    settings.TESTLAB_SLIDES_TEMPLATE = ""
    settings.TESTLAB_SLIDES_TEMPLATE_DEFAULT = tmp_path / "none" / "slides-template.pptx"


@pytest.fixture
def fake_charts(monkeypatch):
    """Chromium is not started: every chart becomes a small PNG."""
    calls = []

    def screenshots(files, document, element_ids, base_url, scale=2):
        calls.append(list(element_ids))
        return {name: png((640, 250)) for name in element_ids}

    monkeypatch.setattr(chromium, "screenshots", screenshots)
    return calls


@pytest.fixture
def report(user):
    """Two final runs in the Final Report; run 001 with pictures and face conditions."""
    run1, result1 = complete_result(user)
    conditions = {"IB Stator": "New", "IB Rotor": "reused", "OB Stator": "", "OB Rotor": "Re-lapped"}
    for component in result1.components.all():
        component.condition = conditions[component.label]
        component.part_no = f"P-{component.label}"
        component.save()
    add_picture(result1, user, ResultImage.Category.CROSS_SECTION, caption="Seal as fitted")
    add_picture(result1, user, ResultImage.Category.PRE_TEST, "IB Stator")
    add_picture(result1, user, ResultImage.Category.POST_TEST, "IB Stator")
    add_picture(result1, user, ResultImage.Category.POST_TEST, "OB Rotor")
    add_picture(result1, user, ResultImage.Category.OTHER, caption="Test rig")
    result1 = finalization.finalize(TestResult.objects.get(pk=result1.pk), user)
    _, result2 = second_run(run1.project, user)
    result2 = finalization.finalize(result2, user)
    return include(run1.project, user, result1, result2)


def deck_of(content):
    return Presentation(io.BytesIO(content))


def texts(slide):
    found = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            found.append(shape.text_frame.text)
        if shape.has_table:
            found += [cell.text for row in shape.table.rows for cell in row.cells]
    return found


def title_of(slide):
    return texts(slide)[0]


def table_rows(slide):
    shape = next(shape for shape in slide.shapes if shape.has_table)
    return [[cell.text for cell in row.cells] for row in shape.table.rows]


def pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == 13]


def test_slides_of_the_final_report_runs(report, fake_charts):
    deck = deck_of(slides.build(report.project))
    assert (deck.slide_width, deck.slide_height) == slides.WIDE
    assert [title_of(s) for s in deck.slides] == [
        report.project.description, "Test matrix", "Parts",
        "Run 001", "Run 001 – Test conditions", "Run 001 – Results",
        "Run 001 – Pictures", "Run 001 – Pictures (continued)",
        "Run 002", "Run 002 – Test conditions", "Run 002 – Results",
    ]  # fmt: skip
    title, matrix, parts, _, conditions, results, photos, more_photos, *_ = deck.slides
    assert [row[0] for row in table_rows(matrix)] == ["Run", "001", "002"]
    assert "Seal as fitted" in texts(conditions)
    assert "Conclusion\nThe faces performed as expected." in texts(conditions)
    assert len(pictures(conditions)) == 1
    rows = table_rows(conditions)
    assert ["Speed", "—", "2980", "rpm"] in rows and ["Test run", "", "", ""] in rows
    assert "Cumulative collected leakage" in texts(results) and len(pictures(results)) == 1
    assert fake_charts == [["chart-0", "chart-1"]]  # one Chromium run for all charts
    # Pre ↔ Post side by side, then the single-stage picture and the other one; the cross section is not repeated.
    assert [t for t in texts(photos) if "·" in t] == ["Pre-test · IB Stator", "Post-test · IB Stator",
                                                       "Post-test · OB Rotor"]  # fmt: skip
    assert "Test rig" in texts(more_photos) and len(pictures(photos) + pictures(more_photos)) == 4


def test_parts_say_new_y_or_n_from_the_face_condition(report, fake_charts):
    parts = deck_of(slides.build(report.project)).slides[2]
    rows = {(row[0], row[1]): (row[4], row[5]) for row in table_rows(parts)[1:]}
    assert rows[("001", "IB Stator")] == ("New", "Y")
    assert rows[("001", "IB Rotor")] == ("reused", "N")
    assert rows[("001", "OB Stator")] == ("—", "")
    assert rows[("001", "OB Rotor")] == ("Re-lapped", "N")


@pytest.mark.parametrize("condition, flag", [("New", "Y"), (" new ", "Y"), ("Reused", "N"), ("Excellent", "N"),
                                             ("", ""), ("—", "")])  # fmt: skip
def test_new_flag(condition, flag):
    assert slides.new_flag(condition) == flag


def test_tester_drawing_on_the_title_slide(report, user, fake_charts):
    out = io.BytesIO()
    Image.new("RGB", (300, 200), "white").save(out, "PDF")
    drawings.add_drawing(report.project, "tester", SimpleUploadedFile("tester.pdf", out.getvalue()), user,
                         "F001842744", "2")  # fmt: skip
    title = deck_of(slides.build(report.project)).slides[0]
    assert "Tester drawing F001842744 rev 2" in texts(title) and len(pictures(title)) == 1


def test_template_layouts_are_used_and_its_slides_removed(report, settings, tmp_path, fake_charts):
    template = Presentation()
    template.slide_width, template.slide_height = slides.WIDE
    template.slides.add_slide(template.slide_layouts[1]).shapes.title.text = "Old slide of the template"
    path = tmp_path / "company.pptx"
    template.save(path)
    settings.TESTLAB_SLIDES_TEMPLATE = str(path)
    deck = deck_of(slides.build(report.project))
    assert "Old slide of the template" not in [title_of(s) for s in deck.slides]
    layouts = [s.slide_layout.name for s in deck.slides]
    assert layouts[:4] == ["Title Slide", "Title Only", "Title Only", "Section Header"]
    assert deck.slides[3].shapes.title.text == "Run 001"
    assert not any(ph.has_text_frame and not ph.text_frame.text for s in deck.slides for ph in s.placeholders)


def test_template_in_the_default_place(report, settings, tmp_path, fake_charts):
    settings.TESTLAB_SLIDES_TEMPLATE_DEFAULT = tmp_path / "slides-template.pptx"
    assert deck_of(slides.build(report.project)).slides[1].slide_layout.name == "Blank"  # none there: plain slides
    template = Presentation()
    template.slide_width, template.slide_height = slides.WIDE
    template.save(tmp_path / "slides-template.pptx")
    assert deck_of(slides.build(report.project)).slides[1].slide_layout.name == "Title Only"


def test_missing_template_is_reported(report, settings, tmp_path):
    settings.TESTLAB_SLIDES_TEMPLATE = str(tmp_path / "gone.pptx")
    with pytest.raises(slides.SlidesError, match="gone.pptx was not found"):
        slides.build(report.project)


def test_download_from_the_final_report_page(client, user, report, fake_charts):
    client.force_login(user)
    page = client.get(reverse("reports:final_report", args=[report.project.pk])).content.decode()
    url = reverse("reports:slides", args=[report.project.pk])
    assert url in page
    response = client.get(url)
    assert response["Content-Type"] == slides.CONTENT_TYPE
    assert response["Content-Disposition"].startswith(f'attachment; filename="{report.project.project_no}-Slides-')
    assert len(deck_of(response.content).slides) == 11


def test_no_runs_included(client, user, report):
    report.runs.all().delete()
    client.force_login(user)
    response = client.get(reverse("reports:slides", args=[report.project.pk]))
    assert response.status_code == 302
    assert [str(m) for m in get_messages(response.wsgi_request)] == [
        "The slides could not be made. No run is included in the Final Report yet: tick the final Results it "
        "covers and save."
    ]


@pytest.mark.chromium
def test_chart_svg_becomes_a_png_with_chromium():
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
           '<rect width="200" height="100" fill="red"/></svg>')
    page = ("<!doctype html><body style='margin:0'>"
            f"<div id='chart-0' style='display:inline-block;line-height:0'>{svg}</div></body>")
    shots = chromium.screenshots({"charts.html": page.encode()}, "charts.html", ["chart-0"], slides.BUNDLE_URL)
    with Image.open(io.BytesIO(shots["chart-0"])) as image:
        assert image.format == "PNG" and image.size == (400, 200)
        assert image.convert("RGB").getpixel((10, 10)) == (255, 0, 0)


def summary_row(kind, quantity, pre, post):
    return {"kind": kind, "quantity": quantity, "pre": {"value": pre}, "post": {"value": post}}


def test_measurement_rows_take_waviness_from_roundpak_and_taper_and_ra_from_mahr():
    ctx = {"measurements": [
        {"component": "IB Stator", "rows": [
            summary_row("roundpak", "Face flatness (circular trace)", "0.412", "0.530"),
            summary_row("roundpak", "Face taper (ROUNDPAK)", "+0.9", "+1.1"),  # not on the slide: Mahr's taper is
            summary_row("mahr", "Face taper (Mahr)", "−0.2100", "+0.0500"),
            summary_row("mahr", "Roughness Ra", "0.051", "0.064"),
            summary_row("dimensional", "Face length", "1.25037", "1.25025"),
        ]},
        {"component": "IB Rotor", "rows": [summary_row("mahr", "Roughness Ra", "0.040", "—")]},
        {"component": "Sleeve", "rows": [summary_row("dimensional", "Dimension", "2.0", "2.0")]},
    ]}  # fmt: skip
    assert slides.measurement_rows(ctx) == [
        ["IB Stator", "0.412", "0.530", "−0.2100", "+0.0500", "0.051", "0.064"],
        ["IB Rotor", "—", "—", "—", "—", "0.040", "—"],
    ]


def test_measurements_slide_after_the_results(report, fake_charts, monkeypatch):
    presented = slides.assembly.presented

    def with_measurements(result):
        ctx, assets = presented(result)
        ctx["measurements"] = [{"component": "IB Stator", "rows": [
            summary_row("mahr", "Face taper (Mahr)", "−0.2100", "+0.0500")]}]  # fmt: skip
        return ctx, assets

    monkeypatch.setattr(slides.assembly, "presented", with_measurements)
    deck = deck_of(slides.build(report.project))
    titles = [title_of(s) for s in deck.slides]
    index = titles.index("Run 002 – Seal-face measurements")
    assert titles[index - 1] == "Run 002 – Results"
    assert table_rows(deck.slides[index]) == [
        ["Component", "Waviness Pre [µm]", "Waviness Post [µm]", "Taper Pre [µm]", "Taper Post [µm]", "Ra Pre [µm]",
         "Ra Post [µm]"],
        ["IB Stator", "—", "—", "−0.2100", "+0.0500", "—", "—"],
    ]  # fmt: skip
