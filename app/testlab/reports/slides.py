"""PowerPoint slides of a project's tests: an editable presentation copy, never a controlled document.

The slides show the runs of the Final Report (its included runs, in its
order) exactly as the Final Report does: from the frozen record of each
run's latest final revision with its Customer report contents
(assembly.presented). Nothing is stored, numbered or issued; the engineer
adjusts the deck in PowerPoint where needed.

    Project   title (with the current tester drawing), test matrix, parts (New = Y/N)
    Per run   divider; test conditions with the conclusion and the seal cross section;
              one slide per DCS plot; leakage and acceptance with the cumulative leakage chart;
              seal-face measurements: waviness (ROUNDPAK), taper and Ra (Mahr), Pre and Post;
              every picture of the Customer Test Run Report, four per slide

Template: data/slides-template.pptx, or the .pptx TESTLAB_SLIDES_TEMPLATE
names; kept out of git (the company template changes every year). Its slides are removed and its layouts
are used: the first layout for the title, a layout named "Section…" for the
run dividers, and a title-only layout for the content. Without a template the
slides are plain widescreen slides. Charts and figures are SVG, which
python-pptx cannot insert: they are turned into PNG with Chromium. A PDF
tester drawing is drawn as a picture of its first page.
"""

import io
import re
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from testlab.documents import chromium
from testlab.documents.render import _BUNDLE_URL as BUNDLE_URL
from testlab.projects import drawings
from testlab.projects.models import ProjectDrawing

from . import assembly

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DASH = "—"
WIDE = (Inches(13.333), Inches(7.5))
MARGIN = Inches(0.4)
GAP = Inches(0.2)
BOTTOM = Inches(0.55)  # kept free for the template's logo and slide number
CAPTION_H = Inches(0.3)
MATRIX_ROWS = 10  # rows per slide; more go on a continuation slide
PARTS_ROWS = 14


class SlidesError(Exception):
    """The slides cannot be made; the message says why."""


def file_name(project):
    return f"{project.project_no}-Slides-{timezone.localdate():%Y%m%d}.pptx"


def build(project):
    """The .pptx bytes of the project's slides."""
    report = assembly.FinalReport.objects.filter(project=project).first()
    results = assembly.included(report)
    if not results:
        raise SlidesError("No run is included in the Final Report yet: tick the final Results it covers and save.")
    runs = []
    for index, result in enumerate(results, start=1):
        ctx, assets = assembly.presented(result)
        runs.append({"ctx": ctx, "assets": assets, "summary": assembly.summary_row(ctx, str(index))})
    charts = _chart_pictures(runs)

    deck = Deck(_template_path())
    _title_slide(deck, project)
    _matrix_slides(deck, [run["summary"] for run in runs])
    _parts_slides(deck, runs)
    for number, run in enumerate(runs):
        _run_slides(deck, run, charts.get(f"chart-{number}"))
    out = io.BytesIO()
    deck.prs.save(out)
    return out.getvalue()


# ---- template ----------------------------------------------------------------------------------


def _template_path():
    name = settings.TESTLAB_SLIDES_TEMPLATE
    if not name:  # the default place: used when a template is there
        default = Path(settings.TESTLAB_SLIDES_TEMPLATE_DEFAULT)
        return default if default.is_file() else None
    path = Path(name)
    if not path.is_file():
        raise SlidesError(f"The slides template {path} was not found (TESTLAB_SLIDES_TEMPLATE).")
    if path.suffix.lower() != ".pptx":
        raise SlidesError(f"The slides template {path} must be a .pptx (save a .potx as a presentation first).")
    return path


def _placeholder_types(layout):
    return [ph.placeholder_format.type for ph in layout.placeholders]


DECORATION = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER}
TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
TEXT_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT}


def _pick_layouts(prs):
    """Title, section and content layout of a template (see the module docstring)."""
    layouts = list(prs.slide_layouts)
    title = layouts[0]
    section = next((layout for layout in layouts if "section" in layout.name.lower()), title)
    title_only = []
    for layout in layouts:
        types = set(_placeholder_types(layout))
        if types & TITLES and types - DECORATION <= TITLES:
            title_only.append(layout)
    named = [layout for layout in title_only if "title only" in layout.name.lower()]
    content = (named or title_only or [layouts[-1]])[0]
    return {"title": title, "section": section, "content": content}


class Deck:
    def __init__(self, template):
        if template is None:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = WIDE
            self.layouts = None
        else:
            self.prs = Presentation(str(template))
            _remove_slides(self.prs)
            self.layouts = _pick_layouts(self.prs)
        self.width, self.height = self.prs.slide_width, self.prs.slide_height

    def slide(self, kind, title, lines=()):
        """A new slide of `kind` (title, section, content) with its title and text lines;
        (slide, (left, top, width, height) of the free area below the title)."""
        if self.layouts is None:
            return self._plain_slide(kind, title, lines)
        slide = self.prs.slides.add_slide(self.layouts[kind])
        title_shape = slide.shapes.title
        others = [ph for ph in slide.placeholders if ph.placeholder_format.type not in TITLES | DECORATION]
        texts = sorted((ph for ph in others if ph.placeholder_format.type in TEXT_PLACEHOLDERS),
                       key=lambda ph: (ph.top or 0, ph.left or 0))  # fmt: skip
        if title_shape is not None:
            title_shape.text_frame.text = title
        for placeholder, line in zip(texts, lines):
            placeholder.text_frame.text = line
        for placeholder in others:  # unused placeholders would show their prompt text in PowerPoint
            if not placeholder.has_text_frame or not placeholder.text_frame.text:
                placeholder.element.getparent().remove(placeholder.element)
        width = self.width - 2 * MARGIN
        if title_shape is None:
            self._text(slide, (MARGIN, Inches(0.25), width, Inches(0.8)), [(title, True, 28)])
            top = Inches(1.15)
        elif kind == "content":
            top = title_shape.top + title_shape.height + GAP
        elif title_shape.top > self.height // 2:  # a title low on the slide: the free area is above it
            top = Inches(0.9)
            return slide, (MARGIN, top, width, title_shape.top - 2 * GAP - top)
        else:
            top = title_shape.top + title_shape.height + Inches(1.2)
        return slide, (MARGIN, top, width, self.height - BOTTOM - top)

    def _plain_slide(self, kind, title, lines):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # Blank
        width = self.width - 2 * MARGIN
        if kind == "content":
            self._text(slide, (MARGIN, Inches(0.25), width, Inches(0.75)), [(title, True, 28)])
            top = Inches(1.1)
            return slide, (MARGIN, top, width, self.height - BOTTOM - top)
        self._text(slide, (MARGIN, Inches(4.0), width, Inches(1.1)), [(title, True, 40)])
        self._text(slide, (MARGIN, Inches(5.15), width, Inches(1.2)), [(line, False, 18) for line in lines])
        return slide, (MARGIN, Inches(0.6), width, Inches(3.2))

    @staticmethod
    def _text(slide, box, paragraphs, anchor=MSO_ANCHOR.TOP):
        shape = slide.shapes.add_textbox(*box)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        for index, (text, bold, size) in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(size)
        return shape


def _remove_slides(prs):
    """Drop the template's own slides; their parts are not written with the deck."""
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


# ---- building blocks ---------------------------------------------------------------------------


def text(slide, box, paragraphs):
    """paragraphs: [(text, bold, size pt)]."""
    return Deck._text(slide, box, paragraphs)


def table(slide, box, header, rows, widths, size=11, groups=()):
    """A table at the top of `box`; `widths` are relative column widths. `groups`: indexes of rows that are
    group headings (one merged, bold cell)."""
    left, top, width, _ = box
    shape = slide.shapes.add_table(len(rows) + 1, len(header), left, top, width, Inches(0.3) * (len(rows) + 1))
    grid = shape.table
    total = sum(widths)
    for index, relative in enumerate(widths):
        grid.columns[index].width = Emu(int(width * relative / total))
    for r, values in enumerate([header, *rows]):
        grid.rows[r].height = Inches(0.28)
        for c, value in enumerate(values):
            _cell(grid.cell(r, c), value, size, bold=r == 0 or (r - 1) in groups)
        if (r - 1) in groups and len(header) > 1:
            grid.cell(r, 0).merge(grid.cell(r, len(header) - 1))
    return shape


def _cell(cell, value, size, bold=False):
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    frame = cell.text_frame
    frame.word_wrap = True
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold


def picture(slide, content, box, caption=""):
    """The picture as large as fits in `box` (centred), with its caption under it."""
    left, top, width, height = box
    if caption:
        height -= CAPTION_H
    with Image.open(io.BytesIO(content)) as image:
        pixels_w, pixels_h = image.size
    scale = min(width / pixels_w, height / pixels_h)
    w, h = int(pixels_w * scale), int(pixels_h * scale)
    x, y = left + (width - w) // 2, top + (height - h) // 2
    slide.shapes.add_picture(io.BytesIO(content), x, y, w, h)
    if caption:
        shape = text(slide, (left, y + h, width, CAPTION_H), [(caption, False, 11)])
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def split(box, fraction, gap=GAP):
    """`box` split into a left part of `fraction` of its width and the rest."""
    left, top, width, height = box
    first = int((width - gap) * fraction)
    return (left, top, first, height), (left + first + gap, top, width - first - gap, height)


def stack(box, fraction, gap=GAP):
    """`box` split into an upper part of `fraction` of its height and the rest."""
    left, top, width, height = box
    first = int((height - gap) * fraction)
    return (left, top, width, first), (left, top + first + gap, width, height - first - gap)


def chunks(items, size):
    return [items[i : i + size] for i in range(0, len(items), size)] or [[]]


def continued(title, index):
    return title if index == 0 else f"{title} (continued)"


# ---- pictures that are not files of the record -------------------------------------------------


def _chart_pictures(runs):
    """PNG of each run's cumulative leakage chart (SVG), all in one Chromium run: {"chart-<n>": PNG}."""
    svgs = {f"chart-{n}": run["ctx"].get("leakage_chart") for n, run in enumerate(runs)}
    svgs = {name: svg for name, svg in svgs.items() if svg}
    if not svgs:
        return {}
    body = "".join(f'<div id="{name}" style="display:inline-block;line-height:0">{svg}</div>'
                   for name, svg in svgs.items())  # fmt: skip
    page = f'<!doctype html><meta charset="utf-8"><body style="margin:0;font-family:Arial,sans-serif">{body}</body>'
    return chromium.screenshots({"charts.html": page.encode()}, "charts.html", svgs, BUNDLE_URL)


# ---- project slides ----------------------------------------------------------------------------


def _title_slide(deck, project):
    lines = [project.project_no, f"{timezone.localdate():%B %Y}"]
    slide, area = deck.slide("title", project.description, lines)
    tester = drawings.current(project, ProjectDrawing.Kind.TESTER)
    if tester is not None:
        _, right = split(area, 0.5)
        picture(slide, drawings.pictures(tester, max_pages=1)[0], right, f"Tester drawing {tester}")


def _matrix_slides(deck, summaries):
    header = ["Run", "TCN", "Dates", "Speed [rpm]", "Medium", "Pressure", "Running [h]", "Avg. dynamic leakage [ml/h]"]
    rows = [[s["run_no"], s["tcn_no"], s["dates"], s["speed"], "\n".join(s["media"]) or DASH,
             "\n".join(s["pressures"]) or DASH, s["running_hours"], "\n".join(s["leakage"]) or DASH]
            for s in summaries]  # fmt: skip
    for index, part in enumerate(chunks(rows, MATRIX_ROWS)):
        slide, area = deck.slide("content", continued("Test matrix", index))
        table(slide, area, header, part, [0.6, 0.8, 1.6, 0.9, 1.8, 1.6, 0.9, 1.8], size=11)


def new_flag(condition):
    """New = Y when the face condition says New, N for any other condition, blank when none is entered."""
    condition = (condition or "").strip()
    if not condition or condition == DASH:
        return ""
    return "Y" if condition.lower() == "new" else "N"


def _parts_slides(deck, runs):
    header = ["Run", "Component", "Part drawing", "Material", "Face condition", "New"]
    rows = []
    for run in runs:
        for c in run["ctx"]["configuration"]["components"]:
            rows.append([run["summary"]["run_no"], c["label"], c["part_no"] or DASH, c["material"] or DASH,
                         c["condition"] or DASH, new_flag(c["condition"])])  # fmt: skip
    for index, part in enumerate(chunks(rows, PARTS_ROWS)):
        slide, area = deck.slide("content", continued("Parts", index))
        table(slide, area, header, part, [0.6, 1.4, 1.6, 2.4, 1.4, 0.5], size=11)


# ---- run slides --------------------------------------------------------------------------------


def _paragraph_lines(paragraphs):
    """The report's paragraphs ([[line, …], …]) as text lines."""
    return [" ".join(lines) for lines in paragraphs if lines]


def _run_slides(deck, run, chart):
    ctx, assets, summary = run["ctx"], run["assets"], run["summary"]
    ident = ctx["identification"]
    name = f"Run {ident['run_no']}"
    deck.slide("section", name, [f"{ident['project_no']} {ident['tcn_no']}", ident.get("dates") or ""])

    # Test conditions, conclusion and the seal cross section.
    slide, area = deck.slide("content", f"{name} – Test conditions")
    left, right = split(area, 0.56)
    header, rows, groups = ["Parameter", "Instructed", "Actual", "Unit"], [], []
    for group in ctx["conditions"]["groups"]:
        groups.append(len(rows))
        rows.append([group["heading"], "", "", ""])
        rows += [[row["parameter"], row["instructed"], row["actual"], row["unit"]] for row in group["rows"]]
    table(slide, left, header, rows, [2.2, 1.6, 1.6, 0.8], size=10, groups=groups)
    cross_sections = ctx["pictures"]["cross_sections"]
    words, image_box = stack(right, 0.45) if cross_sections else (right, None)
    conclusion = _paragraph_lines(ctx["narrative"]["conclusion"]) or [DASH]
    text(slide, words, [("Conclusion", True, 14)] + [(line, False, 12) for line in conclusion])
    if cross_sections:
        first = cross_sections[0]
        picture(slide, assets[first["src"]], image_box, first["caption"])

    # One slide per DCS plot.
    for plot in ctx["plots"]:
        heading = plot["caption"] or f"DCS plot {plot['number']}"
        slide, area = deck.slide("content", f"{name} – {heading}")
        picture(slide, assets[plot["src"]], area, plot["signals"] if not plot["caption"] else "")

    # Running time, leakage and acceptance; the cumulative leakage chart.
    slide, area = deck.slide("content", f"{name} – Results")
    left, right = split(area, 0.5) if chart else (area, None)
    lines = [(f"Shaft running {ident.get('running_hours') or DASH} h · {ident.get('dates') or DASH}", True, 12)]
    text(slide, (left[0], left[1], left[2], Inches(0.4)), lines)
    below = (left[0], left[1] + Inches(0.5), left[2], left[3] - Inches(0.5))
    leakage = []
    for seal in ctx.get("leakage", []):
        criteria = seal.get("criteria", [])
        leakage.append([seal["seal"], seal["total"], seal["average"],
                        "\n".join(c["requirement"] for c in criteria) or DASH,
                        "\n".join(c["label"] for c in criteria) or DASH])  # fmt: skip
    if leakage:
        table(slide, below, ["Seal", "Total [ml]", "Avg. dynamic [ml/h]", "Requirement", "Assessment"], leakage,
              [1.5, 1, 1.3, 1.2, 1], size=11)  # fmt: skip
        used = Inches(0.3) * (len(leakage) + 1) + Inches(0.4)
        below = (below[0], below[1] + used, below[2], below[3] - used)
    acceptance = [[a["criterion"], a["requirement"] or DASH, a["actual"] or DASH, a["label"]]
                  for a in ctx.get("acceptance", [])]  # fmt: skip
    if acceptance:
        table(slide, below, ["Acceptance criterion", "Requirement", "Actual", "Assessment"], acceptance,
              [2.2, 1.3, 1.3, 1], size=10)  # fmt: skip
    if chart:
        picture(slide, chart, right, "Cumulative collected leakage")

    _measurement_slide(deck, name, ctx)

    # Every picture of the Customer report, four per slide (Pre ↔ Post side by side).
    for index, rows in enumerate(chunks(_picture_rows(ctx, skip=cross_sections[:1]), 2)):
        if not rows:
            break
        slide, area = deck.slide("content", continued(f"{name} – Pictures", index))
        for row, row_box in zip(rows, stack(area, 0.5)):
            for cell, cell_box in zip(row, split(row_box, 0.5)):
                if cell is None:
                    continue
                heading, item = cell
                head = text(slide, (cell_box[0], cell_box[1], cell_box[2], CAPTION_H), [(heading, True, 12)])
                head.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                body = (cell_box[0], cell_box[1] + CAPTION_H, cell_box[2], cell_box[3] - CAPTION_H)
                picture(slide, assets[item["src"]], body, item["caption"])


# Slide columns: (heading, kind, summary quantity) of the Pre/Post measurement summary.
MEASUREMENTS = [
    ("Waviness", "roundpak", "Face flatness (circular trace)"),
    ("Taper", "mahr", "Face taper (Mahr)"),
    ("Ra", "mahr", "Roughness Ra"),
]


def measurement_rows(ctx):
    """[component, Pre and Post of each MEASUREMENTS column, …] for the components with any of them."""
    rows = []
    for group in ctx.get("measurements", []):
        found = {(row["kind"], row["quantity"]): row for row in group["rows"]}
        cells = []
        for _, kind, quantity in MEASUREMENTS:
            row = found.get((kind, quantity))
            cells += [row["pre"]["value"], row["post"]["value"]] if row else [DASH, DASH]
        if any(cell != DASH for cell in cells):
            rows.append([group["component"], *cells])
    return rows


def _measurement_slide(deck, name, ctx):
    rows = measurement_rows(ctx)
    if not rows:
        return
    slide, area = deck.slide("content", f"{name} – Seal-face measurements")
    header = ["Component"] + [f"{heading} {stage} [µm]" for heading, _, _ in MEASUREMENTS for stage in ("Pre", "Post")]
    table(slide, area, header, rows, [1.6] + [1] * (len(header) - 1), size=12)
    used = Inches(0.3) * (len(rows) + 1) + GAP
    text(slide, (area[0], area[1] + used, area[2], Inches(0.6)), [
        ("Waviness: ROUNDPAK instrument flatness FLTt of one circular trace. Taper (signed: negative ID high, "
         "positive OD high) and Ra (mean of the left and right scans): Mahr.", False, 11)])  # fmt: skip


def _picture_rows(ctx, skip=()):
    """Rows of two (heading, picture) cells, in the report's order; None leaves a cell empty."""
    pictures, skipped = ctx["pictures"], {p["src"] for p in skip}
    rows = []
    for comparison in pictures["comparisons"]:
        for pair in comparison["pairs"]:
            rows.append([(f"Pre-test · {comparison['label']}", pair["pre"]) if pair["pre"] else None,
                         (f"Post-test · {comparison['label']}", pair["post"]) if pair["post"] else None])  # fmt: skip
    for group in pictures["single_stage"]:
        for row in group["rows"]:
            rows.append([(f"{group['stage']} · {item['label']}", item["picture"]) for item in row] + [None])
    others = [("Seal cross section", p) for p in pictures["cross_sections"] if p["src"] not in skipped]
    others += [(p["caption"] or re.sub(r"\.[^.]+$", "", p["name"]), {**p, "caption": ""}) for p in pictures["other"]]
    rows += [list(pair) + [None] for pair in chunks(others, 2) if pair]
    return [row[:2] for row in rows]
